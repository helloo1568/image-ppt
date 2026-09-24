import json
import sys

import render_deck
from PIL import Image
from pptx import Presentation


def test_render_deck_writes_every_slide_and_comparison(tmp_path, monkeypatch):
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    for number, color in enumerate(("white", "black"), 1):
        path = tmp_path / f"source-{number}.png"
        Image.new("RGB", (160, 90), color).save(path)
    spec = tmp_path / "page-spec.json"
    spec.write_text(json.dumps({"slides": [
        {"id": "s01", "image_status": "approved", "image_file": "source-1.png"},
        {"id": "s02", "image_status": "approved", "image_file": "source-2.png"},
    ]}), encoding="utf-8")

    def fake_renderer(_deck, output, width, height):
        for index in range(1, 3):
            Image.new("RGB", (width, height), "white").save(output / f"{index:03d}.png")

    monkeypatch.setattr(render_deck, "render_powerpoint", fake_renderer)
    monkeypatch.setattr(sys, "argv", ["render_deck.py", str(deck), str(tmp_path / "review"),
                                      "--backend", "powerpoint", "--page-spec", str(spec)])
    render_deck.main()
    report = json.loads((tmp_path / "review/render-report.json").read_text(encoding="utf-8"))
    assert len(report["slides"]) == 2
    assert report["slides"][0]["mean_absolute_difference"] == 0
    assert report["slides"][1]["mean_absolute_difference"] == 255
    assert (tmp_path / "review/review.png").is_file()
    assert all((tmp_path / "review" / f"{index:03d}.png").is_file() for index in (1, 2))
