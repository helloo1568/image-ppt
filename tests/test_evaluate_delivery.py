import json
from pathlib import Path

import audit_page_content
import evaluate_delivery
import pytest
import render_deck
from PIL import Image
from pptx import Presentation


def make_case(tmp_path):
    spec = json.loads((Path(__file__).resolve().parents[1] / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    slide = spec["slides"][0]
    slide["image_status"] = "approved"
    image_path = tmp_path / slide["image_file"]
    image_path.parent.mkdir()
    Image.new("RGB", (160, 90), "white").save(image_path)
    spec_path = tmp_path / "page-spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    rendered = tmp_path / "001.png"
    Image.new("RGB", (160, 90), "white").save(rendered)
    report = {"deck": str(deck), "deck_sha256": render_deck.sha256(deck),
              "page_spec_sha256": render_deck.sha256(spec_path), "backend": "test",
              "slides": [{"slide": 1, "rendered": str(rendered), "rendered_sha256": render_deck.sha256(rendered)}]}
    render_path = tmp_path / "render-report.json"
    render_path.write_text(json.dumps(report), encoding="utf-8")
    review = evaluate_delivery.review_template(spec, report)
    review_path = tmp_path / "visual-review.json"
    review_path.write_text(json.dumps(review), encoding="utf-8")
    observations = audit_page_content.template(spec_path, spec["slides"])
    observations["slides"][0].update(status="complete", observed_text="让数字服务真正走进田间\n以可信、易用的数字服务连接农户与产业资源")
    observation_path = tmp_path / "observations.json"
    observation_path.write_text(json.dumps(observations, ensure_ascii=False), encoding="utf-8")
    return spec_path, render_path, review_path, observation_path, rendered


def test_scorecard_needs_visual_review_and_current_artifacts(tmp_path):
    spec, render, review, observation, rendered = make_case(tmp_path)
    result = evaluate_delivery.evaluate(spec, render, review, [observation])
    assert result["status"] == "incomplete"
    assert result["content"]["issues"] == 0
    reviewed = json.loads(review.read_text(encoding="utf-8"))
    reviewed["slides"][0]["status"] = "pass"
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == "pass"
    Image.new("RGB", (160, 90), "black").save(rendered)
    with pytest.raises(ValueError, match="Rendered slide 1 changed"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


def test_scorecard_rejects_render_without_current_page_spec(tmp_path):
    spec, render, review, observation, _ = make_case(tmp_path)
    report = json.loads(render.read_text(encoding="utf-8"))
    report.pop("page_spec_sha256")
    render.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="lacks the current Page Spec hash"):
        evaluate_delivery.evaluate(spec, render, review, [observation])
