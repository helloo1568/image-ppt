import copy
import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import build_image_ppt
import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "build_image_ppt.py"

COVER_CASES = {
    "01-wide.png": (2000, 1000),
    "02-tall.png": (1000, 2000),
    "03-exact.png": (1920, 1080),
    "04-square.png": (1000, 1000),
}


def make_slides(tmp_path: Path, sizes: dict[str, tuple[int, int]]) -> Path:
    directory = tmp_path / "slides"
    directory.mkdir()
    for name, size in sizes.items():
        Image.new("RGB", size, "#3A6EA5").save(directory / name)
    return directory


def run_build(slides_dir: Path, output: Path, *extra: str) -> dict:
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), str(slides_dir), str(output), *extra],
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout)


@pytest.fixture
def approved_spec(tmp_path):
    spec = json.loads((REPO_ROOT / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    spec["canvas"] = {"width": 400, "height": 300}
    original = spec["slides"][0]
    for element in original["elements"]:
        element.pop("bbox_hint", None)
    spec["slides"] = []
    slides = tmp_path / "slides"
    slides.mkdir()
    # Deliberately oppose natural filename order and leave a superseded draft.
    for page, (name, color) in enumerate((("z-first.png", "red"), ("a-second.png", "blue")), 1):
        slide = copy.deepcopy(original)
        slide.update(id=f"s{page}", page_number=page, image_file=f"slides/{name}", image_status="approved")
        spec["slides"].append(slide)
        Image.new("RGB", (400, 300), color).save(slides / name)
    Image.new("RGB", (400, 300), "green").save(slides / "00-old-draft.png")
    path = tmp_path / "page-spec.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


@pytest.mark.parametrize("dimensions", [(), ("--height", "6")])
def test_page_spec_controls_order_membership_and_aspect_ratio(approved_spec, dimensions):
    output = approved_spec.with_suffix(".pptx")
    result = run_build(approved_spec, output, *dimensions)
    assert result["slides"] == 2
    assert result["images"] == ["z-first.png", "a-second.png"]
    assert result["page_spec"] == str(approved_spec.resolve())
    prs = Presentation(output)
    assert prs.slide_width / prs.slide_height == pytest.approx(4 / 3)
    if dimensions:
        assert result["slide_size_inches"] == [8, 6]
    colors = []
    for slide in prs.slides:
        image = Image.open(BytesIO(pictures(slide)[0].image.blob)).convert("RGB")
        colors.append(image.getpixel((0, 0)))
    assert colors == [(255, 0, 0), (0, 0, 255)]


@pytest.mark.parametrize("mutation", ["unapproved", "missing", "duplicate_image", "wrong_ratio", "unresolved", "dimensions"])
def test_page_spec_build_rejects_invalid_delivery_and_keeps_output(approved_spec, mutation):
    spec = json.loads(approved_spec.read_text(encoding="utf-8"))
    slide = spec["slides"][0]
    extra = []
    if mutation == "unapproved":
        slide["image_status"] = "revision"
    elif mutation == "missing":
        slide["image_file"] = "slides/missing.png"
    elif mutation == "duplicate_image":
        spec["slides"][1]["image_file"] = slide["image_file"]
    elif mutation == "wrong_ratio":
        Image.new("RGB", (300, 400), "red").save(approved_spec.parent / slide["image_file"])
    elif mutation == "unresolved":
        slide["elements"][0].update(confirmation_status="unresolved", confidence=0.5)
    elif mutation == "dimensions":
        extra = ["--width", "16", "--height", "9"]
    approved_spec.write_text(json.dumps(spec), encoding="utf-8")
    output = approved_spec.with_suffix(".pptx")
    output.write_bytes(b"previous approved deck")
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(approved_spec), str(output), *extra],
        cwd=approved_spec.parent, capture_output=True, text=True, check=False,
    )
    assert result.returncode != 0
    assert output.read_bytes() == b"previous approved deck"


def pictures(slide) -> list:
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_natural_sort(tmp_path):
    slides = make_slides(tmp_path, {"2.png": (10, 10), "10.png": (10, 10), "1.png": (10, 10)})
    names = [path.name for path in build_image_ppt.collect_images(slides, (".png", ".jpg", ".jpeg"))]
    assert names == ["1.png", "2.png", "10.png"]


def test_cover_crops_inside_canvas(tmp_path):
    slides = make_slides(tmp_path, COVER_CASES)
    output = tmp_path / "deck.pptx"
    result = run_build(slides, output, "--fit", "cover")
    assert result["slides"] == len(COVER_CASES)
    prs = Presentation(output)
    slide_ratio = prs.slide_width / prs.slide_height
    for slide, (name, (width, height)) in zip(prs.slides, COVER_CASES.items()):
        assert len(pictures(slide)) == 1
        picture = pictures(slide)[0]
        assert picture.name == Path(name).stem
        assert (picture.left, picture.top, picture.width, picture.height) == (
            0,
            0,
            prs.slide_width,
            prs.slide_height,
        )
        image_ratio = width / height
        if image_ratio > slide_ratio:
            expected = (1 - slide_ratio / image_ratio) / 2
            assert abs(picture.crop_left - expected) < 1e-4
            assert abs(picture.crop_right - expected) < 1e-4
            assert picture.crop_top == 0
            assert picture.crop_bottom == 0
        elif image_ratio < slide_ratio:
            expected = (1 - image_ratio / slide_ratio) / 2
            assert abs(picture.crop_top - expected) < 1e-4
            assert abs(picture.crop_bottom - expected) < 1e-4
            assert picture.crop_left == 0
            assert picture.crop_right == 0
        else:
            assert picture.crop_left == 0
            assert picture.crop_right == 0
            assert picture.crop_top == 0
            assert picture.crop_bottom == 0


def test_contain_centers_without_distortion(tmp_path):
    slides = make_slides(tmp_path, COVER_CASES)
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--fit", "contain")
    prs = Presentation(output)
    slide_ratio = prs.slide_width / prs.slide_height
    for slide, (name, (width, height)) in zip(prs.slides, COVER_CASES.items()):
        picture = pictures(slide)[0]
        assert picture.left >= 0 and picture.top >= 0
        assert picture.left + picture.width <= prs.slide_width
        assert picture.top + picture.height <= prs.slide_height
        assert abs(picture.left - (prs.slide_width - picture.width) / 2) <= 2
        assert abs(picture.top - (prs.slide_height - picture.height) / 2) <= 2
        assert abs(picture.width / picture.height - width / height) < 0.01
        if width / height != slide_ratio:
            assert picture.crop_left == 0


def test_stretch_fills_slide(tmp_path):
    slides = make_slides(tmp_path, COVER_CASES)
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--fit", "stretch")
    prs = Presentation(output)
    for slide in prs.slides:
        picture = pictures(slide)[0]
        assert (picture.left, picture.top, picture.width, picture.height) == (
            0,
            0,
            prs.slide_width,
            prs.slide_height,
        )


def test_background_is_slide_fill_without_shapes(tmp_path):
    slides = make_slides(tmp_path, {"01.png": (640, 360)})
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--background", "FF0000")
    prs = Presentation(output)
    slide = prs.slides[0]
    assert len(pictures(slide)) == 1
    assert [shape for shape in slide.shapes if shape.shape_type != MSO_SHAPE_TYPE.PICTURE] == []
    fill = slide.background.fill
    assert fill.type == MSO_FILL.SOLID
    assert fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)


def test_max_width_downscales(tmp_path):
    slides = make_slides(tmp_path, {"01.png": (3000, 1500)})
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--max-width", "1920")
    picture = pictures(Presentation(output).slides[0])[0]
    assert picture.image.size == (1920, 960)


def test_jpeg_quality_reencodes_and_shrinks(tmp_path):
    slides = tmp_path / "slides"
    slides.mkdir()
    # RGB noise: PNG stores it poorly, JPEG compresses it well, so the
    # size assertion is robust across Pillow versions and platforms.
    noise = Image.effect_noise((3000, 1500), 128)
    rgb = Image.merge(
        "RGB",
        [
            noise,
            noise.transpose(Image.Transpose.ROTATE_180),
            noise.transpose(Image.Transpose.FLIP_LEFT_RIGHT),
        ],
    )
    rgb.save(slides / "01.png")
    plain = tmp_path / "plain.pptx"
    run_build(slides, plain)
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--max-width", "1920", "--jpeg-quality", "85")
    picture = pictures(Presentation(output).slides[0])[0]
    assert picture.image.ext == "jpg"
    assert output.stat().st_size < plain.stat().st_size


def test_jpeg_flattens_transparency_onto_background(tmp_path):
    slides = tmp_path / "slides"
    slides.mkdir()
    Image.new("RGBA", (400, 300), (255, 0, 0, 0)).save(slides / "01.png")
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--jpeg-quality", "85", "--background", "3366CC")
    picture = pictures(Presentation(output).slides[0])[0]
    assert picture.image.ext == "jpg"
    embedded = Image.open(BytesIO(picture.image.blob)).convert("RGB")
    pixel = embedded.getpixel((5, 5))
    assert all(abs(channel - expected) <= 4 for channel, expected in zip(pixel, (51, 102, 204)))


def test_exif_orientation_is_applied(tmp_path):
    slides = tmp_path / "slides"
    slides.mkdir()
    exif = Image.Exif()
    exif[0x0112] = 6  # rotate 90 degrees
    Image.new("RGB", (400, 200), "#3A6EA5").save(slides / "01.jpg", exif=exif)
    output = tmp_path / "deck.pptx"
    run_build(slides, output, "--fit", "contain")
    prs = Presentation(output)
    picture = pictures(prs.slides[0])[0]
    assert picture.width / picture.height == pytest.approx(0.5, abs=0.01)


def test_rejects_bad_jpeg_quality(tmp_path):
    slides = make_slides(tmp_path, {"01.png": (640, 360)})
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), str(slides), str(tmp_path / "o.pptx"), "--jpeg-quality", "100"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode != 0
    assert "jpeg-quality" in proc.stderr


def test_json_output_reports_settings(tmp_path):
    slides = make_slides(tmp_path, {"01.png": (640, 360)})
    result = run_build(slides, tmp_path / "deck.pptx", "--max-width", "1280", "--jpeg-quality", "85")
    assert result["max_width"] == 1280
    assert result["jpeg_quality"] == 85
    assert result["images"] == ["01.png"]
