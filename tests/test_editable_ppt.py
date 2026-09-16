import copy
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from audit_editability import audit
from build_editable_ppt import build_deck
from extract_assets import extract
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from scene import load_scene, validate_scene


@pytest.fixture
def scene_file(tmp_path):
    Image.new("RGBA", (200, 100), (10, 100, 200, 120)).save(tmp_path / "asset.png")
    scene = {
        "version": "1.0",
        "canvas": {"width": 1600, "height": 900},
        "slides": [
            {
                "id": "s1",
                "notes": "Source: verified test data",
                "elements": [
                    {
                        "id": "title",
                        "type": "text",
                        "x": 100,
                        "y": 40,
                        "w": 1400,
                        "h": 100,
                        "text": "中文与 English\n增长 25%",
                        "font": "Microsoft YaHei",
                        "font_size": 30,
                    },
                    {
                        "id": "photo",
                        "type": "image",
                        "x": 900,
                        "y": 500,
                        "w": 400,
                        "h": 300,
                        "path": "asset.png",
                    },
                    {
                        "id": "flow",
                        "type": "group",
                        "children": [
                            {
                                "id": "node",
                                "type": "shape",
                                "x": 100,
                                "y": 200,
                                "w": 250,
                                "h": 100,
                                "shape": "roundRect",
                                "fill": "#008877",
                            },
                            {
                                "id": "nested",
                                "type": "group",
                                "children": [
                                    {
                                        "id": "label",
                                        "type": "text",
                                        "x": 110,
                                        "y": 220,
                                        "w": 200,
                                        "h": 60,
                                        "text": "独立节点",
                                        "font_size": 20,
                                    }
                                ],
                            },
                            {
                                "id": "arrow",
                                "type": "line",
                                "x1": 350,
                                "y1": 250,
                                "x2": 550,
                                "y2": 250,
                                "arrow": "end",
                            },
                        ],
                    },
                    {
                        "id": "table",
                        "type": "table",
                        "x": 100,
                        "y": 400,
                        "w": 650,
                        "h": 230,
                        "rows": [["阶段", "数值"], ["第一期", "10"]],
                        "column_widths": [2, 1],
                    },
                    {
                        "id": "chart",
                        "type": "chart",
                        "x": 900,
                        "y": 180,
                        "w": 600,
                        "h": 300,
                        "chart_type": "column",
                        "categories": ["A", "B"],
                        "series": [
                            {"name": "增长", "values": [10, 20], "color": "#008877"}
                        ],
                    },
                ],
            }
        ],
    }
    path = tmp_path / "scene.json"
    path.write_text(json.dumps(scene, ensure_ascii=False), encoding="utf-8")
    return path


def test_native_roundtrip(scene_file):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    report = audit(output, scene_file)
    assert report["errors"] == []
    assert report["warnings"] == []
    assert report["slides"][0]["native_objects"] == 6
    assert report["slides"][0]["movable_raster_objects"] == 1
    prs = Presentation(output)
    assert prs.slides[0].shapes[0].text == "中文与 English\n增长 25%"
    assert (
        prs.slides[0].notes_slide.notes_text_frame.text == "Source: verified test data"
    )
    assert prs.slide_width == Inches(13.333333)
    with ZipFile(output) as package:
        assert any(
            p.startswith("ppt/embeddings/") and p.endswith(".xlsx")
            for p in package.namelist()
        )
        assert b"tailEnd" in package.read("ppt/slides/slide1.xml")


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie", "doughnut"])
def test_chart_types(scene_file, chart_type):
    scene, _ = load_scene(scene_file)
    chart = scene["slides"][0]["elements"][-1]
    chart.update(chart_type=chart_type, data_labels=True, legend=True)
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []


@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
def test_image_fit_preserves_alpha(scene_file, fit):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][1]["fit"] = fit
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    picture = Presentation(output).slides[0].shapes[1]
    import io

    with Image.open(io.BytesIO(picture.image.blob)) as image:
        assert image.getchannel("A").getextrema() == (120, 120)
    if fit == "cover":
        assert picture.crop_left > 0
    elif fit == "contain":
        assert picture.width / picture.height == pytest.approx(2)


@pytest.mark.parametrize(
    "mutation",
    [
        "duplicate",
        "outside",
        "unknown",
        "nan",
        "ragged",
        "values",
        "traversal",
        "missing",
        "zero_line",
        "negative_pie",
    ],
)
def test_invalid_scene_rejected(scene_file, mutation):
    scene, _ = load_scene(scene_file)
    els = scene["slides"][0]["elements"]
    if mutation == "duplicate":
        els.append(copy.deepcopy(els[0]))
    elif mutation == "outside":
        els[0]["x"] = 1500
    elif mutation == "unknown":
        els[0]["typo"] = True
    elif mutation == "nan":
        els[0]["w"] = float("nan")
    elif mutation == "ragged":
        els[-2]["rows"].append(["bad"])
    elif mutation == "values":
        els[-1]["series"][0]["values"] = [10]
    elif mutation == "traversal":
        els[1]["path"] = "../asset.png"
    elif mutation == "missing":
        els[1]["path"] = "missing.png"
    elif mutation == "zero_line":
        els[2]["children"][-1]["x2"] = 350
    elif mutation == "negative_pie":
        els[-1]["chart_type"] = "pie"
        els[-1]["series"][0]["values"] = [-1, 2]
    with pytest.raises(ValueError):
        validate_scene(scene, scene_file.parent)


def test_failed_build_keeps_existing_output(scene_file):
    output = scene_file.with_suffix(".pptx")
    output.write_bytes(b"existing deck")
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][1]["path"] = "missing.png"
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    with pytest.raises(ValueError):
        build_deck(scene_file, output)
    assert output.read_bytes() == b"existing deck"


def test_audit_detects_edits_and_missing_objects(scene_file):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    prs.slides[0].shapes[0].text = "Changed"
    node = prs.slides[0].shapes[1]._element
    node.getparent().remove(node)
    prs.save(output)
    errors = audit(output, scene_file)["errors"]
    assert any("text differs" in e for e in errors)
    assert any("missing object" in e for e in errors)


def test_reference_image_rejected_even_after_reencoding(scene_file):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["source_image"] = "asset.png"
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert any("reference pixels" in e for e in audit(output, scene_file)["errors"])


def test_extract_masks_and_bounds(tmp_path):
    Image.new("RGBA", (100, 60), (250, 0, 0, 128)).save(tmp_path / "source.png")
    Image.new("L", (100, 60), 128).save(tmp_path / "mask.png")
    path = tmp_path / "extract.json"
    spec = {
        "source": "source.png",
        "regions": [{"id": "item", "box": [10.2, 5.3, 30, 20], "mask": "mask.png"}],
    }
    path.write_text(json.dumps(spec), encoding="utf-8")
    result = extract(path, tmp_path / "assets")
    assert result["assets"][0]["source_box"] == [10, 5, 31, 21]
    with Image.open(tmp_path / "assets/item.png") as im:
        assert im.getchannel("A").getextrema() == (64, 64)
    with pytest.raises(ValueError, match="exists"):
        extract(path, tmp_path / "assets")
    spec["regions"][0]["box"] = [90, 50, 30, 20]
    path.write_text(json.dumps(spec), encoding="utf-8")
    with pytest.raises(ValueError, match="outside"):
        extract(path, tmp_path / "bad")
    assert not (tmp_path / "bad").exists()


def test_cli_from_other_working_directory(scene_file, tmp_path):
    script = Path(__file__).resolve().parents[1] / "scripts/build_editable_ppt.py"
    output = tmp_path / "space folder" / "result.pptx"
    result = subprocess.run(
        [sys.executable, str(script), str(scene_file), str(output)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert audit(output, scene_file)["errors"] == []


def test_native_data_can_be_changed_after_export(scene_file):
    from pptx.chart.data import CategoryChartData

    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    slide = prs.slides[0]
    slide.shapes[0].left += Inches(0.2)
    slide.shapes[-2].table.cell(1, 1).text = "99"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("增长", [99, 100])
    slide.shapes[-1].chart.replace_data(data)
    prs.save(output)
    reopened = Presentation(output).slides[0]
    assert list(reopened.shapes[-1].chart.series[0].values) == [99, 100]
    assert reopened.shapes[-2].table.cell(1, 1).text == "99"
    errors = audit(output, scene_file)["errors"]
    assert any("chart data differs" in e for e in errors)
    assert any("table content differs" in e for e in errors)
    assert any("geometry differs" in e for e in errors)
