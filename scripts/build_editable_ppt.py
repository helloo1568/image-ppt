#!/usr/bin/env python3
"""Compile an agent-authored scene JSON into named, editable PowerPoint objects."""

from __future__ import annotations

import argparse
import io
import json
import os
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from scene import asset_path, load_scene

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.CHEVRON,
    "rightArrow": MSO_SHAPE.RIGHT_ARROW,
}
CHARTS = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def rgb(value):
    return RGBColor.from_string(value.lstrip("#"))


def east_asian_font(font, name):
    font.name = name
    for child in list(font._rPr):
        if child.tag.endswith("}ea"):
            font._rPr.remove(child)
    ea = OxmlElement("a:ea")
    ea.set("typeface", name)
    font._rPr.append(ea)


def text_frame(frame, text, e):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }[e.get("valign", "top")]
    for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, attr, Pt(e.get("margin", 0)))
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[e.get("align", "left")]
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = e.get("line_spacing", 1.1)
        font = p.font
        east_asian_font(font, e.get("font", "Arial"))
        font.size = Pt(e.get("font_size", 20))
        font.bold, font.italic = e.get("bold", False), e.get("italic", False)
        font.color.rgb = rgb(e.get("color", "#172B4D"))


def fill_and_line(shape, e):
    fill = e.get("fill")
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if e.get("line"):
        shape.line.color.rgb = rgb(e["line"])
        shape.line.width = Pt(e.get("line_width", 1))
    else:
        shape.line.fill.background()


def add_picture(shapes, e, base, box):
    with Image.open(asset_path(base, e["path"])) as original:
        im = ImageOps.exif_transpose(original).convert("RGBA")
        iw, ih = im.size
        data = io.BytesIO()
        im.save(data, format="PNG")
    data.seek(0)
    x, y, w, h = box
    image_ratio, box_ratio = iw / ih, w / h
    fit = e.get("fit", "contain")
    if fit == "contain":
        scale = min(w / iw, h / ih)
        nw, nh = round(iw * scale), round(ih * scale)
        x, y = x + (w - nw) // 2, y + (h - nh) // 2
        w, h = nw, nh
    picture = shapes.add_picture(data, x, y, w, h)
    if fit == "cover":
        if image_ratio > box_ratio:
            picture.crop_left = picture.crop_right = (1 - box_ratio / image_ratio) / 2
        else:
            picture.crop_top = picture.crop_bottom = (1 - image_ratio / box_ratio) / 2
    return picture


def add_element(shapes, e, base, sx, sy):
    kind = e["type"]
    box = (
        tuple(
            round(e[k] * scale)
            for k, scale in zip(("x", "y", "w", "h"), (sx, sy, sx, sy))
        )
        if kind not in ("group", "line")
        else None
    )
    if kind == "group":
        shape = shapes.add_group_shape()
        for child in e["children"]:
            add_element(shape.shapes, child, base, sx, sy)
    elif kind == "text":
        shape = shapes.add_textbox(*box)
        text_frame(shape.text_frame, e["text"], e)
    elif kind == "shape":
        shape = shapes.add_shape(SHAPES[e["shape"]], *box)
        fill_and_line(shape, e)
    elif kind == "image":
        shape = add_picture(shapes, e, base, box)
    elif kind == "line":
        shape = shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            round(e["x1"] * sx),
            round(e["y1"] * sy),
            round(e["x2"] * sx),
            round(e["y2"] * sy),
        )
        shape.line.color.rgb = rgb(e.get("color", "#172B4D"))
        shape.line.width = Pt(e.get("line_width", 1.5))
        arrow = e.get("arrow", "none")
        for tag in (
            ["a:tailEnd"]
            if arrow == "end"
            else ["a:headEnd", "a:tailEnd"]
            if arrow == "both"
            else []
        ):
            node = OxmlElement(tag)
            node.set("type", "triangle")
            shape._element.spPr.get_or_add_ln().append(node)
    elif kind == "table":
        rows = e["rows"]
        shape = shapes.add_table(len(rows), len(rows[0]), *box)
        table = shape.table
        table.first_row = e.get("header", True)
        table.horz_banding = False
        if "column_widths" in e:
            total = sum(e["column_widths"])
            widths = [round(box[2] * weight / total) for weight in e["column_widths"]]
            widths[-1] += box[2] - sum(widths)
            for col, width in zip(table.columns, widths):
                col.width = width
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                header = r == 0 and e.get("header", True)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(
                    e.get("header_fill", "#172B4D")
                    if header
                    else e.get("fill", "#FFFFFF")
                )
                style = {
                    **e,
                    "color": e.get("header_color", "#FFFFFF")
                    if header
                    else e.get("color", "#172B4D"),
                    "bold": header or e.get("bold", False),
                    "margin": e.get("margin", 5),
                }
                text_frame(cell.text_frame, value, style)
    elif kind == "chart":
        data = CategoryChartData()
        data.categories = e["categories"]
        for series in e["series"]:
            data.add_series(series["name"], series["values"])
        shape = shapes.add_chart(CHARTS[e["chart_type"]], *box, data)
        chart = shape.chart
        chart.has_title = False
        chart.has_legend = e.get("legend", len(e["series"]) > 1)
        east_asian_font(chart.font, e.get("font", "Arial"))
        chart.font.size = Pt(e.get("font_size", 14))
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.has_data_labels = e.get("data_labels", False)
        if plot.has_data_labels:
            plot.data_labels.position = (
                XL_LABEL_POSITION.BEST_FIT
                if e["chart_type"] in ("pie", "doughnut")
                else XL_LABEL_POSITION.OUTSIDE_END
                if e["chart_type"] in ("bar", "column")
                else XL_LABEL_POSITION.ABOVE
            )
            plot.data_labels.number_format = e.get("number_format", "General")
            plot.data_labels.number_format_is_linked = False
            east_asian_font(plot.data_labels.font, e.get("font", "Arial"))
        if e["chart_type"] not in ("pie", "doughnut"):
            for axis in (chart.category_axis, chart.value_axis):
                east_asian_font(axis.tick_labels.font, e.get("font", "Arial"))
        for series, spec in zip(chart.series, e["series"]):
            if spec.get("color"):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = rgb(spec["color"])
                series.format.line.color.rgb = rgb(spec["color"])
    else:
        raise ValueError(f"Unsupported element type: {kind}")
    shape.name = f"{e['id']} | {e.get('name', kind)}"
    if kind in ("shape", "line"):
        # Override the default Office theme's effectRef (otherwise shapes gain shadows).
        shape._element.spPr.append(OxmlElement("a:effectLst"))
    if "rotation" in e:
        shape.rotation = e["rotation"]
    return shape


def build_deck(scene_path: Path, output: Path) -> dict:
    scene_path, output = scene_path.resolve(), output.resolve()
    if output.suffix.lower() != ".pptx":
        raise ValueError("Output must have .pptx extension")
    scene, warnings = load_scene(scene_path)
    canvas = scene["canvas"]
    prs = Presentation()
    prs.slide_width = Inches(canvas.get("width_inches", 13.333333))
    prs.slide_height = round(prs.slide_width * canvas["height"] / canvas["width"])
    prs.core_properties.title = scene.get("title", output.stem)
    sx, sy = prs.slide_width / canvas["width"], prs.slide_height / canvas["height"]
    for spec in scene["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(spec.get("background", "#FFFFFF"))
        for e in spec["elements"]:
            add_element(slide.shapes, e, scene_path.parent, sx, sy)
        slide.notes_slide.notes_text_frame.text = spec.get("notes", "")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Save/reopen before replacing an existing deck; failed builds leave it intact.
    fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=output.parent)
    os.close(fd)
    try:
        prs.save(tmp)
        verified = Presentation(tmp)
        if len(verified.slides) != len(scene["slides"]):
            raise RuntimeError("Slide count changed during export")
        os.replace(tmp, output)
    finally:
        Path(tmp).unlink(missing_ok=True)
    return {
        "output": str(output),
        "slides": len(prs.slides),
        "warnings": warnings,
        "visual_review": "required: render and inspect in a presentation renderer",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("scene", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    try:
        result = build_deck(args.scene, args.output)
    except (ValueError, OSError) as error:
        parser.exit(2, f"build_editable_ppt: {error}\n")
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
